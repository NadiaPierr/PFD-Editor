#version 0.0.10
from io import BytesIO
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import customtkinter as ctk
import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pdf2docx import Converter
from PIL import Image, ImageDraw, ImageOps, ImageTk
from tkinterdnd2 import DND_FILES, TkinterDnD


class PDFEditorApp(ctk.CTk, TkinterDnD.DnDWrapper):
    def __init__(self):
        super().__init__()
        self.TkdndVersion = TkinterDnD._require(self)

        self.title("PDF Editor")
        self.geometry("1100x720")
        self.minsize(860, 560)
        self.after(0, self.open_maximized)
        self.configure(fg_color=("#f4f6f8", "#111317"))

        self.pdf_document = None
        self.pdf_path = None
        self.current_page_index = 0
        self.preview_image = None
        self.preview_base_image = None
        self.preview_images = []
        self.page_layouts = []
        self.preview_scale = 1.0
        self.zoom = 1.0
        self.resize_job = None
        self.extract_mode = False
        self.selection_start = None
        self.selection_end = None
        self.selection_page_index = None
        self.selection_anchor_word_index = None
        self.selected_word_indices = []

        ctk.set_appearance_mode("System")
        ctk.set_default_color_theme("blue")

        self._build_ui()

    def open_maximized(self):
        try:
            self.state("zoomed")
        except Exception:
            self.attributes("-fullscreen", True)
        self.after(100, self.center_preview_placeholder)

    def _build_ui(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        sidebar = ctk.CTkFrame(
            self,
            width=240,
            corner_radius=0,
            fg_color=("#ffffff", "#171a20"),
        )
        sidebar.grid(row=0, column=0, sticky="nsw")
        sidebar.grid_propagate(False)
        sidebar.grid_columnconfigure(0, weight=1)
        sidebar.grid_rowconfigure(15, weight=1)

        title = ctk.CTkLabel(
            sidebar,
            text="PDF Editor",
            font=ctk.CTkFont(size=26, weight="bold"),
        )
        title.grid(row=0, column=0, padx=18, pady=(28, 6), sticky="w")

        subtitle = ctk.CTkLabel(
            sidebar,
            text="Open, preview, and convert PDF files.",
            text_color=("#667085", "#9aa4b2"),
            anchor="w",
            justify="left",
            wraplength=190,
        )
        subtitle.grid(row=1, column=0, padx=18, pady=(0, 24), sticky="ew")

        self.open_button = ctk.CTkButton(
            sidebar,
            text="Open PDF",
            height=44,
            corner_radius=8,
            command=self.open_pdf,
        )
        self.open_button.grid(row=2, column=0, padx=18, pady=(0, 14), sticky="ew")

        drop_hint = ctk.CTkLabel(
            sidebar,
            text="You can also drag a PDF into the preview area.",
            text_color=("#667085", "#9aa4b2"),
            anchor="w",
            justify="left",
            wraplength=190,
        )
        drop_hint.grid(row=3, column=0, padx=18, pady=(0, 22), sticky="ew")

        details_label = ctk.CTkLabel(
            sidebar,
            text="Document",
            font=ctk.CTkFont(size=14, weight="bold"),
            anchor="w",
            text_color=("#344054", "#e4e7ec"),
        )
        details_label.grid(row=4, column=0, padx=18, pady=(0, 8), sticky="ew")

        self.file_label = ctk.CTkLabel(
            sidebar,
            text="No file selected.",
            fg_color=("#f2f4f7", "#20242c"),
            corner_radius=8,
            anchor="w",
            justify="left",
            height=46,
            wraplength=180,
            text_color=("#344054", "#e4e7ec"),
        )
        self.file_label.grid(row=5, column=0, padx=18, pady=(0, 10), sticky="ew")

        self.page_label = ctk.CTkLabel(
            sidebar,
            text="Pages: -",
            text_color=("#667085", "#9aa4b2"),
            anchor="w",
            justify="left",
        )
        self.page_label.grid(row=6, column=0, padx=18, pady=(0, 24), sticky="ew")

        theme_label = ctk.CTkLabel(
            sidebar,
            text="Theme",
            font=ctk.CTkFont(size=14, weight="bold"),
            anchor="w",
            text_color=("#344054", "#e4e7ec"),
        )
        theme_label.grid(row=7, column=0, padx=18, pady=(0, 8), sticky="ew")

        self.theme_selector = ctk.CTkSegmentedButton(
            sidebar,
            values=["System", "Light", "Dark"],
            command=self.change_theme,
        )
        self.theme_selector.grid(row=8, column=0, padx=18, pady=(0, 24), sticky="ew")
        self.theme_selector.set("System")

        actions_label = ctk.CTkLabel(
            sidebar,
            text="Actions",
            font=ctk.CTkFont(size=14, weight="bold"),
            anchor="w",
            text_color=("#344054", "#e4e7ec"),
        )
        actions_label.grid(row=9, column=0, padx=18, pady=(0, 8), sticky="ew")

        self.convert_button = ctk.CTkButton(
            sidebar,
            text="Convert to Word",
            height=42,
            corner_radius=8,
            fg_color=("#344054", "#344054"),
            hover_color=("#1d2939", "#475467"),
            command=self.convert_pdf_to_word,
            state="disabled",
        )
        self.convert_button.grid(row=10, column=0, padx=18, pady=(0, 12), sticky="new")

        self.merge_button = ctk.CTkButton(
            sidebar,
            text="Merge PDFs",
            height=42,
            corner_radius=8,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=self.merge_pdf_files,
        )
        self.merge_button.grid(row=11, column=0, padx=18, pady=(0, 12), sticky="new")

        self.extract_button = ctk.CTkButton(
            sidebar,
            text="Extract Text",
            height=42,
            corner_radius=8,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=self.toggle_extract_mode,
            state="disabled",
        )
        self.extract_button.grid(row=12, column=0, padx=18, pady=(0, 12), sticky="new")

        self.split_button = ctk.CTkButton(
            sidebar,
            text="Split PDF",
            height=42,
            corner_radius=8,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=self.show_split_pdf_dialog,
            state="disabled",
        )
        self.split_button.grid(row=13, column=0, padx=18, pady=(0, 12), sticky="new")

        self.images_to_pdf_button = ctk.CTkButton(
            sidebar,
            text="Images to PDF",
            height=42,
            corner_radius=8,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=self.convert_images_to_pdf,
        )
        self.images_to_pdf_button.grid(row=14, column=0, padx=18, pady=(0, 12), sticky="new")

        self.status_label = ctk.CTkLabel(
            sidebar,
            text="",
            text_color=("#667085", "#9aa4b2"),
            anchor="w",
            justify="left",
            wraplength=190,
        )
        self.status_label.grid(row=15, column=0, padx=18, pady=(0, 24), sticky="sew")

        preview_area = ctk.CTkFrame(self, fg_color=("#f4f6f8", "#111317"), corner_radius=0)
        preview_area.grid(row=0, column=1, padx=0, pady=0, sticky="nsew")
        preview_area.grid_columnconfigure(0, weight=1)
        preview_area.grid_rowconfigure(1, weight=1)

        top_bar = ctk.CTkFrame(preview_area, fg_color="transparent")
        top_bar.grid(row=0, column=0, columnspan=2, padx=28, pady=(24, 12), sticky="ew")
        top_bar.grid_columnconfigure(0, weight=1)

        self.page_number_label = ctk.CTkLabel(
            top_bar,
            text="No document",
            font=ctk.CTkFont(size=15, weight="bold"),
            text_color=("#344054", "#e4e7ec"),
            anchor="w",
        )
        self.page_number_label.grid(row=0, column=0, sticky="w")

        zoom_controls = ctk.CTkFrame(top_bar, fg_color="transparent")
        zoom_controls.grid(row=0, column=1, sticky="e")
        zoom_controls.grid_columnconfigure((0, 1, 2), weight=0)

        self.zoom_out_button = ctk.CTkButton(
            zoom_controls,
            text="-",
            width=34,
            height=30,
            command=self.zoom_out,
            state="disabled",
        )
        self.zoom_out_button.grid(row=0, column=0, padx=(0, 6))

        self.zoom_label = ctk.CTkLabel(
            zoom_controls,
            text="100%",
            width=58,
            text_color=("#344054", "#e4e7ec"),
        )
        self.zoom_label.grid(row=0, column=1, padx=(0, 6))

        self.zoom_in_button = ctk.CTkButton(
            zoom_controls,
            text="+",
            width=34,
            height=30,
            command=self.zoom_in,
            state="disabled",
        )
        self.zoom_in_button.grid(row=0, column=2)

        self.zoom_reset_button = ctk.CTkButton(
            zoom_controls,
            text="Fit",
            width=44,
            height=30,
            command=self.reset_zoom,
            state="disabled",
        )
        self.zoom_reset_button.grid(row=0, column=3, padx=(6, 0))

        self.preview_frame = ctk.CTkFrame(
            preview_area,
            fg_color=("#e9edf2", "#161a20"),
            corner_radius=12,
        )
        self.preview_frame.grid(row=1, column=0, padx=(28, 10), pady=(0, 28), sticky="nsew")
        self.preview_frame.grid_columnconfigure(0, weight=1)
        self.preview_frame.grid_rowconfigure(0, weight=1)
        self.preview_frame.bind("<Configure>", self.schedule_preview_rerender)

        self.preview_canvas = tk.Canvas(
            self.preview_frame,
            bg="#e9edf2",
            bd=0,
            highlightthickness=0,
        )
        self.preview_canvas.grid(row=0, column=0, sticky="nsew")
        self.preview_canvas.create_text(
            0,
            0,
            text="Drop a PDF here\nor use Open PDF",
            fill="#667085",
            font=("Arial", 18, "bold"),
            justify="center",
            tags=("placeholder",),
        )
        self.after(0, self.center_preview_placeholder)
        self.preview_canvas.bind("<ButtonPress-1>", self.start_text_selection)
        self.preview_canvas.bind("<B1-Motion>", self.update_text_selection)
        self.preview_canvas.bind("<ButtonRelease-1>", self.finish_text_selection)
        self.preview_canvas.bind("<Configure>", self.on_preview_canvas_configure)

        self.enable_drag_and_drop()

        self.page_scrollbar = ctk.CTkScrollbar(
            preview_area,
            orientation="vertical",
            command=self.on_page_scroll,
            width=16,
        )
        self.page_scrollbar.grid(row=1, column=1, padx=(0, 18), pady=(0, 28), sticky="ns")
        self.page_scrollbar.set(0, 1)
        self.preview_canvas.configure(yscrollcommand=self.on_canvas_scroll)

        self.bind("<MouseWheel>", self.on_mouse_wheel)

    def enable_drag_and_drop(self):
        for widget in (self, self.preview_frame, self.preview_canvas):
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", self.on_file_drop)

    def change_theme(self, theme):
        ctk.set_appearance_mode(theme)

    def on_file_drop(self, event):
        file_paths = self.tk.splitlist(event.data)
        if not file_paths:
            return

        dropped_path = Path(file_paths[0])
        if dropped_path.suffix.lower() != ".pdf":
            messagebox.showwarning("Invalid file", "Drop a PDF file to open it.")
            return

        self.load_pdf(dropped_path)

    def open_pdf(self):
        file_path = filedialog.askopenfilename(
            title="Select PDF",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )

        if not file_path:
            return

        self.load_pdf(Path(file_path))

    def load_pdf(self, file_path):
        try:
            document = fitz.open(str(file_path))
        except Exception as exc:
            messagebox.showerror("Error", f"Could not open the PDF file.\n\n{exc}")
            return

        if document.page_count == 0:
            document.close()
            messagebox.showwarning("Empty PDF", "This PDF does not contain any pages.")
            return

        if self.pdf_document is not None:
            self.pdf_document.close()

        self.pdf_document = document
        self.pdf_path = Path(file_path)
        self.current_page_index = 0

        self.file_label.configure(text=f"File: {self.pdf_path.name}")
        self.status_label.configure(text="")
        self.convert_button.configure(state="normal")
        self.extract_button.configure(state="normal")
        self.split_button.configure(state="normal")
        self.preview_base_image = None
        self.exit_extract_mode(update_status=False)
        self.zoom = 1.0
        self.set_zoom_controls_state("normal")
        self.render_current_page()

    def convert_pdf_to_word(self):
        if self.pdf_path is None:
            messagebox.showwarning("No PDF", "Open a PDF file first.")
            return

        initial_file = f"{self.pdf_path.stem}.docx"
        output_path = filedialog.asksaveasfilename(
            title="Save Word file",
            defaultextension=".docx",
            initialfile=initial_file,
            filetypes=[("Word document", "*.docx"), ("All files", "*.*")],
        )

        if not output_path:
            return

        self.convert_button.configure(state="disabled")
        self.status_label.configure(text="Converting PDF to Word...")
        self.update_idletasks()

        converter = None
        try:
            converter = Converter(str(self.pdf_path))
            converter.convert(
                output_path,
                start=0,
                end=None,
                **self.high_accuracy_conversion_settings(),
            )
            self.optimize_word_tables_for_cell_fit(output_path)
            self.optimize_word_symbols(output_path)
        except Exception as exc:
            messagebox.showerror("Conversion error", f"Could not convert the PDF file.\n\n{exc}")
            self.status_label.configure(text="Conversion failed.")
            return
        finally:
            if converter is not None:
                converter.close()

            self.convert_button.configure(state="normal")

        self.status_label.configure(text=f"Saved: {Path(output_path).name}")
        messagebox.showinfo("Conversion complete", "The Word file was saved successfully.")

    def merge_pdf_files(self):
        file_paths = filedialog.askopenfilenames(
            title="Select PDFs to merge",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )

        if not file_paths:
            return

        selected_paths = [Path(file_path) for file_path in file_paths]
        if len(selected_paths) < 2:
            messagebox.showwarning("Select more PDFs", "Select at least two PDF files to merge.")
            return

        initial_directory = selected_paths[0].parent
        output_path = filedialog.asksaveasfilename(
            title="Save merged PDF",
            defaultextension=".pdf",
            initialdir=initial_directory,
            initialfile="merged.pdf",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )

        if not output_path:
            return

        output_path = Path(output_path)
        if output_path.resolve() in {path.resolve() for path in selected_paths}:
            messagebox.showwarning(
                "Invalid output",
                "Save the merged PDF with a different name from the selected files.",
            )
            return

        self.merge_button.configure(state="disabled")
        self.convert_button.configure(state="disabled")
        self.status_label.configure(text="Merging PDF files...")
        self.update_idletasks()

        merged_document = fitz.open()
        try:
            for selected_path in selected_paths:
                with fitz.open(str(selected_path)) as source_document:
                    if source_document.page_count == 0:
                        continue

                    merged_document.insert_pdf(source_document)

            if merged_document.page_count == 0:
                messagebox.showwarning("Empty merge", "No pages were found in the selected PDFs.")
                self.status_label.configure(text="Merge cancelled.")
                return

            merged_document.save(str(output_path), garbage=4, deflate=True)
        except Exception as exc:
            messagebox.showerror("Merge error", f"Could not merge the PDF files.\n\n{exc}")
            self.status_label.configure(text="Merge failed.")
            return
        finally:
            merged_document.close()
            self.merge_button.configure(state="normal")
            self.convert_button.configure(state="normal" if self.pdf_path is not None else "disabled")

        self.status_label.configure(text=f"Merged: {output_path.name}")
        messagebox.showinfo("Merge complete", "The merged PDF was saved successfully.")
        self.load_pdf(output_path)

    def convert_images_to_pdf(self):
        image_paths = filedialog.askopenfilenames(
            title="Select images",
            filetypes=[
                ("Image files", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"),
                ("PNG files", "*.png"),
                ("JPEG files", "*.jpg *.jpeg"),
                ("All files", "*.*"),
            ],
        )

        if not image_paths:
            return

        selected_paths = [Path(image_path) for image_path in image_paths]
        output_path = filedialog.asksaveasfilename(
            title="Save PDF",
            defaultextension=".pdf",
            initialfile="images.pdf",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )

        if not output_path:
            return

        output_path = Path(output_path)
        self.images_to_pdf_button.configure(state="disabled")
        self.status_label.configure(text="Converting images to PDF...")
        self.update_idletasks()

        pdf_document = fitz.open()
        try:
            for image_path in selected_paths:
                image_bytes, width, height = self.prepare_image_for_pdf(image_path)
                page = pdf_document.new_page(width=width, height=height)
                page.insert_image(page.rect, stream=image_bytes)

            pdf_document.save(str(output_path), garbage=4, deflate=True)
        except Exception as exc:
            messagebox.showerror("Image conversion error", f"Could not create the PDF file.\n\n{exc}")
            self.status_label.configure(text="Image conversion failed.")
            return
        finally:
            pdf_document.close()
            self.images_to_pdf_button.configure(state="normal")

        self.status_label.configure(text=f"Saved: {output_path.name}")
        messagebox.showinfo("Conversion complete", "The image PDF was saved successfully.")
        self.load_pdf(output_path)

    def prepare_image_for_pdf(self, image_path):
        with Image.open(image_path) as image:
            image = ImageOps.exif_transpose(image)

            if image.width <= 0 or image.height <= 0:
                raise ValueError(f"Invalid image size: {image_path.name}")

            if image.mode in ("RGBA", "LA") or "transparency" in image.info:
                alpha = image.convert("RGBA").getchannel("A")
                background = Image.new("RGB", image.size, "white")
                background.paste(image.convert("RGBA"), mask=alpha)
                image = background
            elif image.mode != "RGB":
                image = image.convert("RGB")

            image_buffer = BytesIO()
            image.save(image_buffer, format="PNG")
            return image_buffer.getvalue(), image.width, image.height

    def show_split_pdf_dialog(self):
        if self.pdf_document is None or self.pdf_path is None:
            messagebox.showwarning("No PDF", "Open a PDF file first.")
            return

        total_pages = self.pdf_document.page_count
        window = ctk.CTkToplevel(self)
        window.title("Split PDF")
        window.geometry("460x230")
        window.minsize(420, 220)
        window.transient(self)
        window.grab_set()
        window.grid_columnconfigure(0, weight=1)

        title = ctk.CTkLabel(
            window,
            text="Select pages to split",
            font=ctk.CTkFont(size=18, weight="bold"),
            anchor="w",
        )
        title.grid(row=0, column=0, padx=18, pady=(18, 6), sticky="ew")

        hint = ctk.CTkLabel(
            window,
            text=f"Total pages: {total_pages}. Example: 1, 3, 5-8",
            text_color=("#667085", "#9aa4b2"),
            anchor="w",
        )
        hint.grid(row=1, column=0, padx=18, pady=(0, 10), sticky="ew")

        pages_entry = ctk.CTkEntry(
            window,
            placeholder_text="Pages or ranges",
            height=38,
        )
        pages_entry.grid(row=2, column=0, padx=18, pady=(0, 14), sticky="ew")
        pages_entry.focus_set()

        actions = ctk.CTkFrame(window, fg_color="transparent")
        actions.grid(row=3, column=0, padx=18, pady=(0, 18), sticky="ew")
        actions.grid_columnconfigure(0, weight=1)

        cancel_button = ctk.CTkButton(
            actions,
            text="Cancel",
            width=110,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=window.destroy,
        )
        cancel_button.grid(row=0, column=1, padx=(8, 0), sticky="e")

        split_button = ctk.CTkButton(
            actions,
            text="Split",
            width=110,
            command=lambda: self.split_pdf_pages(pages_entry.get(), window),
        )
        split_button.grid(row=0, column=2, padx=(8, 0), sticky="e")

        window.bind("<Return>", lambda _event: self.split_pdf_pages(pages_entry.get(), window))
        window.lift()

    def split_pdf_pages(self, page_selection, dialog):
        if self.pdf_document is None or self.pdf_path is None:
            messagebox.showwarning("No PDF", "Open a PDF file first.")
            return

        try:
            page_indices = self.parse_page_selection(page_selection, self.pdf_document.page_count)
        except ValueError as exc:
            messagebox.showwarning("Invalid pages", str(exc))
            return

        initial_file = f"{self.pdf_path.stem}_split.pdf"
        output_path = filedialog.asksaveasfilename(
            title="Save split PDF",
            defaultextension=".pdf",
            initialfile=initial_file,
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )

        if not output_path:
            return

        output_path = Path(output_path)
        if output_path.resolve() == self.pdf_path.resolve():
            messagebox.showwarning("Invalid output", "Save the split PDF with a different name.")
            return

        self.split_button.configure(state="disabled")
        self.status_label.configure(text="Splitting PDF...")
        self.update_idletasks()

        split_document = fitz.open()
        try:
            for page_index in page_indices:
                split_document.insert_pdf(
                    self.pdf_document,
                    from_page=page_index,
                    to_page=page_index,
                )

            split_document.save(str(output_path), garbage=4, deflate=True)
        except Exception as exc:
            messagebox.showerror("Split error", f"Could not split the PDF file.\n\n{exc}")
            self.status_label.configure(text="Split failed.")
            return
        finally:
            split_document.close()
            self.split_button.configure(state="normal")

        dialog.destroy()
        self.status_label.configure(text=f"Split: {output_path.name}")
        messagebox.showinfo("Split complete", "The split PDF was saved successfully.")
        self.load_pdf(output_path)

    def parse_page_selection(self, page_selection, total_pages):
        page_selection = page_selection.strip()
        if not page_selection:
            raise ValueError("Enter at least one page or page range.")

        selected_pages = []
        seen_pages = set()

        for part in page_selection.split(","):
            part = part.strip()
            if not part:
                continue

            if "-" in part:
                range_parts = [value.strip() for value in part.split("-", 1)]
                if len(range_parts) != 2 or not all(value.isdigit() for value in range_parts):
                    raise ValueError("Use page numbers like 1, 3, 5-8.")

                start_page, end_page = (int(value) for value in range_parts)
                if start_page > end_page:
                    raise ValueError("Page ranges must go from smaller to larger page numbers.")

                page_numbers = range(start_page, end_page + 1)
            else:
                if not part.isdigit():
                    raise ValueError("Use page numbers like 1, 3, 5-8.")

                page_numbers = [int(part)]

            for page_number in page_numbers:
                if page_number < 1 or page_number > total_pages:
                    raise ValueError(f"Pages must be between 1 and {total_pages}.")

                page_index = page_number - 1
                if page_index not in seen_pages:
                    selected_pages.append(page_index)
                    seen_pages.add(page_index)

        if not selected_pages:
            raise ValueError("Enter at least one valid page.")

        return selected_pages

    def toggle_extract_mode(self):
        if self.pdf_document is None:
            messagebox.showwarning("No PDF", "Open a PDF file first.")
            return

        if self.extract_mode:
            self.exit_extract_mode()
            return

        self.extract_mode = True
        self.selection_start = None
        self.selection_end = None
        self.selection_page_index = None
        self.selection_anchor_word_index = None
        self.selected_word_indices = []
        self.set_extract_button_active(True)
        self.preview_canvas.configure(cursor="xterm")
        self.status_label.configure(text="Drag across text to extract it.")

    def exit_extract_mode(self, update_status=True):
        self.extract_mode = False
        self.selection_start = None
        self.selection_end = None
        self.selection_page_index = None
        self.selection_anchor_word_index = None
        self.selected_word_indices = []

        if hasattr(self, "extract_button"):
            self.set_extract_button_active(False)

        if hasattr(self, "preview_canvas"):
            self.preview_canvas.configure(cursor="")
            self.preview_canvas.delete("selection")

        if update_status:
            self.status_label.configure(text="")

    def set_extract_button_active(self, active):
        if active:
            self.extract_button.configure(
                text="Cancel Extract",
                fg_color=("#16a34a", "#22c55e"),
                hover_color=("#15803d", "#16a34a"),
            )
        else:
            self.extract_button.configure(
                text="Extract Text",
                fg_color=("#475467", "#475467"),
                hover_color=("#344054", "#667085"),
            )

    def start_text_selection(self, event):
        if not self.extract_mode or self.pdf_document is None:
            return

        position = self.get_preview_image_position(event)
        if position is None:
            return

        self.selection_start = position
        self.selection_end = position
        self.selection_page_index = position[0]
        self.selection_anchor_word_index = self.find_word_index_at_position(position)
        self.selected_word_indices = []
        if self.selection_anchor_word_index is not None:
            self.selected_word_indices = [self.selection_anchor_word_index]

        self.draw_text_selection_overlay()

    def update_text_selection(self, event):
        if not self.extract_mode or self.selection_start is None:
            return

        position = self.get_preview_image_position(event, clamp=True)
        if position is None:
            return

        self.selection_end = position
        self.update_selected_words(position)
        self.draw_text_selection_overlay()

    def finish_text_selection(self, event):
        if not self.extract_mode or self.selection_start is None:
            return

        position = self.get_preview_image_position(event, clamp=True)
        if position is not None:
            self.selection_end = position
            self.update_selected_words(position)
            self.draw_text_selection_overlay()

        if not self.selected_word_indices:
            self.selection_start = None
            self.selection_end = None
            self.preview_canvas.delete("selection")
            return

        try:
            extracted_text = self.extract_text_from_selected_words()
        except Exception as exc:
            messagebox.showerror("Extract error", f"Could not extract text.\n\n{exc}")
            self.status_label.configure(text="Text extraction failed.")
            return

        self.status_label.configure(text="Text extracted.")
        self.show_extracted_text_window(extracted_text)
        self.clear_text_selection_state(clear_overlay=False)

    def get_preview_image_position(self, event, clamp=False):
        if not self.page_layouts:
            return None

        canvas_x = self.preview_canvas.canvasx(event.x)
        canvas_y = self.preview_canvas.canvasy(event.y)
        layout = self.page_layout_at(canvas_x, canvas_y)

        if layout is None and clamp and self.selection_page_index is not None:
            layout = self.layout_for_page(self.selection_page_index)

        if layout is None:
            return None

        x = canvas_x - layout["x"]
        y = canvas_y - layout["y"]
        image_width = layout["width"]
        image_height = layout["height"]

        if clamp:
            x = max(0, min(x, image_width - 1))
            y = max(0, min(y, image_height - 1))
            return (layout["page_index"], x, y)

        if x < 0 or y < 0 or x >= image_width or y >= image_height:
            return None

        return (layout["page_index"], x, y)

    def normalized_selection_rect(self):
        if self.selection_start is None or self.selection_end is None:
            return None

        page0, x0, y0 = self.selection_start
        page1, x1, y1 = self.selection_end
        if page0 != page1:
            return None

        left, right = sorted((x0, x1))
        top, bottom = sorted((y0, y1))

        if right - left < 4 or bottom - top < 4:
            return None

        return (left, top, right, bottom)

    def update_selected_words(self, position):
        if self.selection_page_index is None or position[0] != self.selection_page_index:
            return

        words = self.get_current_page_words(self.selection_page_index)
        if not words:
            self.selected_word_indices = []
            return

        current_word_index = self.find_word_index_at_position(position)
        if self.selection_anchor_word_index is not None and current_word_index is not None:
            start = min(self.selection_anchor_word_index, current_word_index)
            end = max(self.selection_anchor_word_index, current_word_index)
            self.selected_word_indices = list(range(start, end + 1))
            return

        selection_rect = self.normalized_selection_rect()
        if selection_rect is None:
            self.selected_word_indices = []
            return

        selected_indices = []
        for index, word in enumerate(words):
            word_rect = self.image_rect_for_word(word, self.selection_page_index)
            if self.selection_overlaps_word(fitz.Rect(selection_rect), word_rect):
                selected_indices.append(index)

        self.selected_word_indices = selected_indices

    def draw_text_selection_overlay(self):
        self.preview_canvas.delete("selection")

        if not self.selected_word_indices or self.selection_page_index is None:
            return

        layout = self.layout_for_page(self.selection_page_index)
        if layout is None:
            return

        words = self.get_current_page_words(self.selection_page_index)

        for word_index in self.selected_word_indices:
            word_rect = self.padded_word_rect(
                self.image_rect_for_word(words[word_index], self.selection_page_index),
                x_padding=1.5,
                y_padding=1.0,
            )
            self.preview_canvas.create_rectangle(
                layout["x"] + word_rect.x0,
                layout["y"] + word_rect.y0,
                layout["x"] + word_rect.x1,
                layout["y"] + word_rect.y1,
                fill="#2f80ed",
                outline="#2f80ed",
                stipple="gray25",
                tags=("selection",),
            )

    def find_word_index_at_position(self, position):
        page_index = position[0]
        words = self.get_current_page_words(page_index)
        if not words:
            return None

        point = fitz.Point(position[1], position[2])
        matching_indices = []

        for index, word in enumerate(words):
            word_rect = self.image_rect_for_word(word, page_index)
            padded_rect = self.padded_word_rect(word_rect)
            if padded_rect.contains(point):
                matching_indices.append((index, self.distance_to_rect_center(point, word_rect)))

        if matching_indices:
            matching_indices.sort(key=lambda match: match[1])
            return matching_indices[0][0]

        nearest_word = self.find_nearest_word_on_same_line(point, words, page_index)
        if nearest_word is not None:
            return nearest_word

        return None

    def find_nearest_word_on_same_line(self, point, words, page_index):
        candidates = []

        for index, word in enumerate(words):
            word_rect = self.image_rect_for_word(word, page_index)
            line_padding = max(word_rect.height * 0.65, 6)
            line_rect = fitz.Rect(
                word_rect.x0,
                word_rect.y0 - line_padding,
                word_rect.x1,
                word_rect.y1 + line_padding,
            )

            if line_rect.y0 <= point.y <= line_rect.y1:
                horizontal_distance = self.horizontal_distance_to_rect(point, word_rect)
                if horizontal_distance <= max(word_rect.width * 0.55, 18):
                    candidates.append((index, horizontal_distance))

        if not candidates:
            return None

        candidates.sort(key=lambda candidate: candidate[1])
        return candidates[0][0]

    def padded_word_rect(self, word_rect, x_padding=None, y_padding=None):
        if x_padding is None:
            x_padding = max(word_rect.width * 0.12, 2)
        if y_padding is None:
            y_padding = max(word_rect.height * 0.35, 3)

        return fitz.Rect(
            word_rect.x0 - x_padding,
            word_rect.y0 - y_padding,
            word_rect.x1 + x_padding,
            word_rect.y1 + y_padding,
        )

    def distance_to_rect_center(self, point, rect):
        center_x = (rect.x0 + rect.x1) / 2
        center_y = (rect.y0 + rect.y1) / 2
        return abs(point.x - center_x) + abs(point.y - center_y)

    def horizontal_distance_to_rect(self, point, rect):
        if rect.x0 <= point.x <= rect.x1:
            return 0

        return min(abs(point.x - rect.x0), abs(point.x - rect.x1))

    def selection_overlaps_word(self, selection_rect, word_rect):
        padded_rect = self.padded_word_rect(word_rect, x_padding=1.5, y_padding=2.5)
        intersection = selection_rect & padded_rect
        if intersection.is_empty:
            return False

        vertical_overlap = intersection.height / max(padded_rect.height, 1)
        horizontal_overlap = intersection.width / max(padded_rect.width, 1)
        return vertical_overlap >= 0.35 and horizontal_overlap >= 0.08

    def get_current_page_words(self, page_index=None):
        if self.pdf_document is None:
            return []

        if page_index is None:
            page_index = self.current_page_index

        page = self.pdf_document.load_page(page_index)
        words = page.get_text("words")
        return sorted(words, key=lambda word: (word[5], word[6], word[7]))

    def image_rect_for_word(self, word, page_index=None):
        if page_index is None:
            page_index = self.current_page_index

        layout = self.layout_for_page(page_index)
        if layout is None:
            return fitz.Rect()

        page = self.pdf_document.load_page(page_index)
        scale_x = layout["width"] / page.rect.width
        scale_y = layout["height"] / page.rect.height
        return fitz.Rect(
            word[0] * scale_x,
            word[1] * scale_y,
            word[2] * scale_x,
            word[3] * scale_y,
        )

    def extract_text_from_selected_words(self):
        if self.selection_page_index is None:
            return ""

        words = self.get_current_page_words(self.selection_page_index)
        selected_words = [words[index] for index in self.selected_word_indices]
        selected_words.sort(key=lambda word: (word[5], word[6], word[7]))

        lines = []
        current_line_key = None
        current_line_words = []

        for word in selected_words:
            line_key = (word[5], word[6])
            if current_line_key is not None and line_key != current_line_key:
                lines.append(" ".join(current_line_words))
                current_line_words = []

            current_line_key = line_key
            current_line_words.append(word[4])

        if current_line_words:
            lines.append(" ".join(current_line_words))

        return "\n".join(lines).strip()

    def extract_text_from_selection(self, selection_rect):
        page = self.pdf_document.load_page(self.current_page_index)
        image_width, image_height = self.preview_base_image.size
        scale_x = image_width / page.rect.width
        scale_y = image_height / page.rect.height
        left, top, right, bottom = selection_rect
        page_rect = fitz.Rect(
            left / scale_x,
            top / scale_y,
            right / scale_x,
            bottom / scale_y,
        )

        text = page.get_textbox(page_rect).strip()
        if text:
            return text

        return self.extract_words_from_rect(page, page_rect)

    def extract_words_from_rect(self, page, page_rect):
        words = page.get_text("words")
        selected_words = [
            word
            for word in words
            if fitz.Rect(word[:4]).intersects(page_rect)
        ]
        selected_words.sort(key=lambda word: (word[5], word[6], word[7]))

        lines = []
        current_line_key = None
        current_line_words = []

        for word in selected_words:
            line_key = (word[5], word[6])
            if current_line_key is not None and line_key != current_line_key:
                lines.append(" ".join(current_line_words))
                current_line_words = []

            current_line_key = line_key
            current_line_words.append(word[4])

        if current_line_words:
            lines.append(" ".join(current_line_words))

        return "\n".join(lines).strip()

    def show_extracted_text_window(self, text):
        window = ctk.CTkToplevel(self)
        window.title("Extracted Text")
        window.geometry("620x420")
        window.minsize(420, 280)
        window.transient(self)
        window.protocol("WM_DELETE_WINDOW", lambda: self.close_extracted_text_window(window))
        window.bind("<Escape>", lambda _event: self.close_extracted_text_window(window))
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(0, weight=1)

        text_box = ctk.CTkTextbox(window, wrap="word")
        text_box.grid(row=0, column=0, columnspan=3, padx=16, pady=(16, 10), sticky="nsew")
        text_box.insert("1.0", text if text else "No text found in the selected area.")

        copy_button = ctk.CTkButton(
            window,
            text="Copy",
            width=110,
            command=lambda: self.copy_extracted_text(text_box.get("1.0", "end-1c")),
        )
        copy_button.grid(row=1, column=0, padx=(16, 8), pady=(0, 16), sticky="w")

        save_button = ctk.CTkButton(
            window,
            text="Save TXT",
            width=110,
            command=lambda: self.save_extracted_text(text_box.get("1.0", "end-1c")),
        )
        save_button.grid(row=1, column=1, padx=8, pady=(0, 16), sticky="w")

        close_button = ctk.CTkButton(
            window,
            text="Close",
            width=110,
            fg_color=("#475467", "#475467"),
            hover_color=("#344054", "#667085"),
            command=lambda: self.close_extracted_text_window(window),
        )
        close_button.grid(row=1, column=2, padx=(8, 16), pady=(0, 16), sticky="e")
        window.focus_force()
        window.lift()

    def close_extracted_text_window(self, window):
        if not window.winfo_exists():
            return

        self.clear_text_selection_state(clear_overlay=True)

        try:
            window.grab_release()
        except Exception:
            pass

        window.destroy()

    def clear_text_selection_state(self, clear_overlay=True):
        self.selection_start = None
        self.selection_end = None
        self.selection_page_index = None
        self.selection_anchor_word_index = None
        self.selected_word_indices = []

        if clear_overlay and hasattr(self, "preview_canvas"):
            self.preview_canvas.delete("selection")

    def copy_extracted_text(self, text):
        self.clipboard_clear()
        self.clipboard_append(text)
        self.status_label.configure(text="Extracted text copied.")

    def save_extracted_text(self, text):
        if self.pdf_path is None:
            initial_file = "extracted_text.txt"
        else:
            initial_file = f"{self.pdf_path.stem}_extracted_page_{self.current_page_index + 1}.txt"

        output_path = filedialog.asksaveasfilename(
            title="Save extracted text",
            defaultextension=".txt",
            initialfile=initial_file,
            filetypes=[("Text file", "*.txt"), ("All files", "*.*")],
        )

        if not output_path:
            return

        Path(output_path).write_text(text, encoding="utf-8")
        self.status_label.configure(text=f"Saved: {Path(output_path).name}")

    def high_accuracy_conversion_settings(self):
        return {
            "ignore_page_error": True,
            "multi_processing": False,
            "parse_lattice_table": True,
            "parse_stream_table": True,
            "extract_stream_table": False,
            "min_section_height": 20.0,
            "connected_border_tolerance": 1.2,
            "min_border_clearance": 1.0,
            "max_border_width": 8.0,
            "shape_min_dimension": 1.0,
            "min_svg_gap_dx": 8.0,
            "min_svg_gap_dy": 1.5,
            "min_svg_w": 1.0,
            "min_svg_h": 1.0,
            "list_not_table": True,
            "line_overlap_threshold": 0.9,
            "line_break_width_ratio": 0.5,
            "line_break_free_space_ratio": 0.1,
            "line_separate_threshold": 5.0,
            "new_paragraph_free_space_ratio": 0.82,
            "max_line_spacing_ratio": 1.35,
            "lines_left_aligned_threshold": 1.2,
            "lines_right_aligned_threshold": 1.2,
            "lines_center_aligned_threshold": 2.0,
            "float_image_ignorable_gap": 1.0,
            "clip_image_res_ratio": 6.0,
            "delete_end_line_hyphen": True,
        }

    def optimize_word_tables_for_cell_fit(self, docx_path):
        document = Document(docx_path)

        for table in self.iter_all_tables(document):
            table.autofit = False
            table.allow_autofit = False
            self.set_table_layout(table, "fixed")
            self.set_table_cell_spacing(table, 0)

            for row in table.rows:
                for cell in row.cells:
                    self.set_cell_margins(cell, top=24, start=36, bottom=24, end=36)

                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_before = Pt(0)
                        paragraph.paragraph_format.space_after = Pt(0)

        self.remove_trailing_empty_paragraphs(document)
        document.save(docx_path)

    def iter_all_tables(self, parent):
        for table in parent.tables:
            yield table

            for row in table.rows:
                for cell in row.cells:
                    yield from self.iter_all_tables(cell)

    def set_table_layout(self, table, layout_type):
        table_properties = table._tbl.tblPr
        table_layout = table_properties.first_child_found_in("w:tblLayout")

        if table_layout is None:
            table_layout = OxmlElement("w:tblLayout")
            table_properties.append(table_layout)

        table_layout.set(qn("w:type"), layout_type)

    def set_table_cell_spacing(self, table, spacing):
        table_properties = table._tbl.tblPr
        cell_spacing = table_properties.first_child_found_in("w:tblCellSpacing")

        if cell_spacing is None:
            cell_spacing = OxmlElement("w:tblCellSpacing")
            table_properties.append(cell_spacing)

        cell_spacing.set(qn("w:w"), str(spacing))
        cell_spacing.set(qn("w:type"), "dxa")

    def set_cell_margins(self, cell, top=0, start=0, bottom=0, end=0):
        table_cell_properties = cell._tc.get_or_add_tcPr()
        table_cell_margins = table_cell_properties.first_child_found_in("w:tcMar")

        if table_cell_margins is None:
            table_cell_margins = OxmlElement("w:tcMar")
            table_cell_properties.append(table_cell_margins)

        for margin_name, margin_value in {
            "top": top,
            "start": start,
            "bottom": bottom,
            "end": end,
        }.items():
            margin_element = table_cell_margins.find(qn(f"w:{margin_name}"))
            if margin_element is None:
                margin_element = OxmlElement(f"w:{margin_name}")
                table_cell_margins.append(margin_element)

            margin_element.set(qn("w:w"), str(margin_value))
            margin_element.set(qn("w:type"), "dxa")

    def remove_trailing_empty_paragraphs(self, document):
        body = document._body._element

        if not len(body):
            return

        last_element = body[-1]
        if last_element.tag == qn("w:p") and self.is_empty_paragraph_element(last_element):
            section_properties = self.get_paragraph_section_properties(last_element)
            if section_properties is not None:
                section_properties.getparent().remove(section_properties)
                body.remove(last_element)
                body.append(section_properties)

        while len(body):
            last_index = len(body) - 1
            if body[last_index].tag == qn("w:sectPr"):
                last_index -= 1

            if last_index < 0:
                break

            candidate = body[last_index]
            if candidate.tag != qn("w:p") or not self.is_empty_paragraph_element(candidate):
                break

            if self.get_paragraph_section_properties(candidate) is not None:
                break

            body.remove(candidate)

    def get_paragraph_section_properties(self, paragraph_element):
        paragraph_properties = paragraph_element.find(qn("w:pPr"))
        if paragraph_properties is None:
            return None

        return paragraph_properties.find(qn("w:sectPr"))

    def is_empty_paragraph_element(self, paragraph_element):
        text = "".join(node.text or "" for node in paragraph_element.iter(qn("w:t")))
        drawings = list(paragraph_element.iter(qn("w:drawing")))
        pictures = list(paragraph_element.iter(qn("w:pict")))

        return not text.strip() and not drawings and not pictures

    def optimize_word_symbols(self, docx_path):
        document = Document(docx_path)

        for paragraph in self.iter_all_paragraphs(document):
            paragraph_has_math = False

            for run in paragraph.runs:
                if not run.text or not self.contains_symbol(run.text):
                    continue

                font_name = self.symbol_font_for_text(run.text)
                if font_name == "Cambria Math":
                    paragraph_has_math = True

                run.font.name = font_name
                run_element = run._element
                run_properties = run_element.get_or_add_rPr()
                run_fonts = run_properties.rFonts

                if run_fonts is None:
                    run_fonts = OxmlElement("w:rFonts")
                    run_properties.append(run_fonts)

                run_fonts.set(qn("w:ascii"), font_name)
                run_fonts.set(qn("w:hAnsi"), font_name)
                run_fonts.set(qn("w:eastAsia"), font_name)
                run_fonts.set(qn("w:cs"), font_name)

            if paragraph_has_math:
                paragraph.paragraph_format.line_spacing = 1.25
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)

        document.save(docx_path)

    def iter_all_paragraphs(self, parent):
        for paragraph in parent.paragraphs:
            yield paragraph

        for table in parent.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from self.iter_all_paragraphs(cell)

    def contains_symbol(self, text):
        return any(self.is_symbol_character(character) for character in text)

    def is_symbol_character(self, character):
        codepoint = ord(character)
        return (
            0x0370 <= codepoint <= 0x03FF
            or 0x2000 <= codepoint <= 0x206F
            or 0x2070 <= codepoint <= 0x209F
            or 0x20A0 <= codepoint <= 0x20CF
            or 0x2100 <= codepoint <= 0x214F
            or 0x2190 <= codepoint <= 0x21FF
            or 0x2200 <= codepoint <= 0x22FF
            or 0x2300 <= codepoint <= 0x23FF
            or 0x2460 <= codepoint <= 0x24FF
            or 0x2500 <= codepoint <= 0x25FF
            or 0x2600 <= codepoint <= 0x27BF
        )

    def symbol_font_for_text(self, text):
        if any(self.is_math_symbol(character) for character in text):
            return "Cambria Math"

        return "Segoe UI Symbol"

    def is_math_symbol(self, character):
        codepoint = ord(character)
        math_symbols = "\u00b1\u00d7\u00f7\u2264\u2265\u2260\u2248\u221a\u221e"
        math_symbols += "\u2211\u220f\u222b\u2202\u2206\u2207\u03c0"
        return (
            0x2070 <= codepoint <= 0x209F
            or 0x2100 <= codepoint <= 0x214F
            or 0x2200 <= codepoint <= 0x22FF
            or character in math_symbols
        )

    def render_current_page(self):
        if self.pdf_document is None:
            return

        total_pages = self.pdf_document.page_count
        self.current_page_index = max(0, min(self.current_page_index, total_pages - 1))
        current_scroll = self.preview_canvas.yview()[0] if self.page_layouts else 0

        try:
            self.selection_start = None
            self.selection_end = None
            self.selection_page_index = None
            self.selection_anchor_word_index = None
            self.selected_word_indices = []
            self.render_document_pages()
            self.preview_canvas.yview_moveto(current_scroll)
        except Exception as exc:
            self.preview_image = None
            self.preview_base_image = None
            self.preview_images = []
            self.page_layouts = []
            self.preview_canvas.delete("all")
            self.preview_canvas.create_text(
                380,
                260,
                text=f"Could not load the PDF preview.\n{exc}",
                fill="#d92d20",
                font=("Arial", 14, "bold"),
                justify="center",
            )

        self.update_current_page_from_scroll()
        self.update_zoom_controls()

    def render_document_pages(self):
        self.preview_canvas.delete("all")
        self.preview_images = []
        self.page_layouts = []

        if self.pdf_document is None:
            return

        canvas_width = self.preview_canvas.winfo_width()
        if canvas_width <= 100:
            canvas_width = self.preview_frame.winfo_width()
        if canvas_width <= 100:
            canvas_width = 760

        max_page_width = max(page.rect.width for page in self.pdf_document)
        available_width = max(canvas_width - 2, 240)
        self.preview_scale = min(available_width / max_page_width, 2.0) * self.zoom

        y = 0
        max_rendered_width = 0
        page_gap = 14

        for page_index in range(self.pdf_document.page_count):
            page = self.pdf_document.load_page(page_index)
            page_image = self.render_page_image(page, self.preview_scale)
            photo_image = ImageTk.PhotoImage(page_image)
            self.preview_images.append(photo_image)

            x = max((canvas_width - page_image.width) // 2, 0)
            self.preview_canvas.create_image(x, y, image=photo_image, anchor="nw")
            self.preview_canvas.create_rectangle(
                x,
                y,
                x + page_image.width - 1,
                y + page_image.height - 1,
                outline="#c2cad6",
                width=1,
            )
            self.page_layouts.append(
                {
                    "page_index": page_index,
                    "x": x,
                    "y": y,
                    "width": page_image.width,
                    "height": page_image.height,
                }
            )
            max_rendered_width = max(max_rendered_width, x + page_image.width)
            y += page_image.height + page_gap

        scroll_height = max(y - page_gap, self.preview_canvas.winfo_height())
        self.preview_canvas.configure(scrollregion=(0, 0, max_rendered_width, scroll_height))

    def render_page_image(self, page, scale):
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        page_image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        draw = ImageDraw.Draw(page_image)
        draw.rectangle(
            (0, 0, page_image.width - 1, page_image.height - 1),
            outline="#c2cad6",
            width=1,
        )
        return page_image

    def page_layout_at(self, canvas_x, canvas_y):
        for layout in self.page_layouts:
            if (
                layout["x"] <= canvas_x <= layout["x"] + layout["width"]
                and layout["y"] <= canvas_y <= layout["y"] + layout["height"]
            ):
                return layout

        return None

    def layout_for_page(self, page_index):
        for layout in self.page_layouts:
            if layout["page_index"] == page_index:
                return layout

        return None

    def render_browser_style_page(self, page):
        frame_width = self.preview_frame.winfo_width()
        frame_height = self.preview_frame.winfo_height()
        if frame_width <= 100:
            frame_width = 760
        if frame_height <= 100:
            frame_height = 920

        viewport_padding = 0
        shadow_size = 0
        available_width = max(frame_width - viewport_padding * 2 - shadow_size, 240)
        available_height = max(frame_height - viewport_padding * 2 - shadow_size, 320)
        scale = min(
            available_width / page.rect.width,
            available_height / page.rect.height,
            2.0,
        ) * self.zoom

        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        page_image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)

        viewport_width = page_image.width
        viewport_height = page_image.height
        viewport = Image.new("RGB", (viewport_width, viewport_height), "white")

        page_x = 0
        page_y = 0
        draw = ImageDraw.Draw(viewport)
        viewport.paste(page_image, (page_x, page_y))
        draw.rectangle(
            (page_x, page_y, page_x + page_image.width - 1, page_y + page_image.height - 1),
            outline="#c2cad6",
            width=1,
        )
        return viewport

    def schedule_preview_rerender(self, _event=None):
        if self.pdf_document is None:
            self.center_preview_placeholder()
            return

        if self.resize_job is not None:
            self.after_cancel(self.resize_job)

        self.resize_job = self.after(180, self.rerender_after_resize)

    def rerender_after_resize(self):
        self.resize_job = None
        self.render_current_page()

    def center_preview_placeholder(self):
        if not hasattr(self, "preview_canvas"):
            return

        canvas_width = max(self.preview_canvas.winfo_width(), 1)
        canvas_height = max(self.preview_canvas.winfo_height(), 1)
        center_x = self.preview_canvas.canvasx(canvas_width / 2)
        center_y = self.preview_canvas.canvasy(canvas_height / 2)
        self.preview_canvas.coords("placeholder", center_x, center_y)

    def on_preview_canvas_configure(self, _event=None):
        if self.pdf_document is None:
            self.center_preview_placeholder()

    def on_canvas_scroll(self, first, last):
        self.page_scrollbar.set(first, last)
        self.update_current_page_from_scroll()

    def on_page_scroll(self, *args):
        if self.pdf_document is None or not self.page_layouts:
            return

        self.preview_canvas.yview(*args)
        self.update_current_page_from_scroll()

    def on_mouse_wheel(self, event):
        if self.pdf_document is None:
            return

        if event.state & 0x0004:
            if event.delta > 0:
                self.zoom_in()
            else:
                self.zoom_out()
            return

        scroll_steps = -6 if event.delta > 0 else 6
        self.preview_canvas.yview_scroll(scroll_steps, "units")
        self.update_current_page_from_scroll()
        return "break"

    def go_to_page(self, page_index):
        if self.pdf_document is None or not self.page_layouts:
            return

        total_pages = self.pdf_document.page_count
        page_index = max(0, min(page_index, total_pages - 1))
        layout = self.layout_for_page(page_index)
        if layout is None:
            return

        self.current_page_index = page_index
        scrollregion = self.preview_canvas.bbox("all")
        if scrollregion is None:
            return

        total_height = max(scrollregion[3] - scrollregion[1], 1)
        self.preview_canvas.yview_moveto(layout["y"] / total_height)
        self.update_current_page_from_scroll()

    def update_current_page_from_scroll(self):
        if self.pdf_document is None or not self.page_layouts:
            self.page_number_label.configure(text="No document")
            self.page_label.configure(text="Pages: -")
            return

        viewport_top = self.preview_canvas.canvasy(0)
        viewport_middle = viewport_top + max(self.preview_canvas.winfo_height(), 1) / 2
        current_layout = min(
            self.page_layouts,
            key=lambda layout: abs((layout["y"] + layout["height"] / 2) - viewport_middle),
        )
        self.current_page_index = current_layout["page_index"]
        current_page = self.current_page_index + 1
        total_pages = self.pdf_document.page_count
        self.page_number_label.configure(text=f"Page {current_page}")
        self.page_label.configure(text=f"Page {current_page} of {total_pages}")

    def zoom_in(self):
        self.set_zoom(self.zoom + 0.1)

    def zoom_out(self):
        self.set_zoom(self.zoom - 0.1)

    def reset_zoom(self):
        self.set_zoom(1.0)

    def set_zoom(self, zoom):
        if self.pdf_document is None:
            return

        self.zoom = max(0.5, min(2.5, zoom))
        self.render_current_page()

    def update_zoom_controls(self):
        self.zoom_label.configure(text=f"{round(self.zoom * 100)}%")

        if self.pdf_document is None:
            self.set_zoom_controls_state("disabled")
            return

        self.zoom_out_button.configure(state="normal" if self.zoom > 0.5 else "disabled")
        self.zoom_in_button.configure(state="normal" if self.zoom < 2.5 else "disabled")
        self.zoom_reset_button.configure(state="normal")

    def set_zoom_controls_state(self, state):
        self.zoom_out_button.configure(state=state)
        self.zoom_in_button.configure(state=state)
        self.zoom_reset_button.configure(state=state)

    def destroy(self):
        if self.resize_job is not None:
            self.after_cancel(self.resize_job)
            self.resize_job = None

        if self.pdf_document is not None:
            self.pdf_document.close()

        super().destroy()


if __name__ == "__main__":
    app = PDFEditorApp()
    app.mainloop()
