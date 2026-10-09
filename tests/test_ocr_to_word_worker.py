"""Tests for the Word-document helpers used by the OCR worker."""

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from types import SimpleNamespace

from src.ocr_to_word_worker import (
    add_digital_page_text,
    add_html_table,
    add_layout_blocks,
    normalized_text_lines,
    order_positioned_items,
)


def test_add_html_table_creates_a_simple_table_with_cell_text():
    document = Document()

    add_html_table(
        document,
        "<table><tr><td>Name</td><td>Score</td></tr>"
        "<tr><td>Ada</td><td>10</td></tr></table>",
    )

    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["Name", "Score"],
        ["Ada", "10"],
    ]


def test_add_html_table_supports_headers_line_breaks_and_merged_cells():
    document = Document()

    add_html_table(
        document,
        "<table>"
        "<tr><th colspan='2'>Results</th></tr>"
        "<tr><td rowspan='2'>Ada<br>Smith</td><td>First</td></tr>"
        "<tr><td>Second</td></tr>"
        "</table>",
    )

    table = document.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Results"
    assert table.cell(0, 1).text == "Results"
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    assert table.cell(1, 0).text == "Ada\nSmith"
    assert table.cell(2, 0).text == "Ada\nSmith"
    assert table.cell(1, 1).text == "First"
    assert table.cell(2, 1).text == "Second"


def test_add_html_table_falls_back_to_a_paragraph_when_no_rows_are_parsed():
    document = Document()
    html = "Not a table at all"

    add_html_table(document, html)

    assert not document.tables
    assert [paragraph.text for paragraph in document.paragraphs] == [html]


def test_add_html_table_handles_empty_html_without_creating_a_table():
    document = Document()

    add_html_table(document, "")

    assert not document.tables
    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].text == ""


def test_add_layout_blocks_sorts_and_formats_mixed_ocr_content():
    """OCR blocks must follow their visual reading order in the DOCX."""
    document = Document()
    page_image = Image.new("RGB", (400, 300), "white")
    section = document.sections[0]

    # Deliberately provide the blocks out of order, as can happen in OCR output.
    result = {
        "parsing_res_list": [
            {"block_label": "image", "block_bbox": [20, 220, 120, 280]},
            {"block_label": "table", "block_content": "<table><tr><td>Cell</td></tr></table>", "block_bbox": [20, 160, 180, 200]},
            {"block_label": "text", "block_content": "Opening paragraph", "block_bbox": [20, 80, 250, 110]},
            {"block_label": "title", "block_content": "Document title", "block_bbox": [20, 20, 250, 50]},
        ]
    }

    add_layout_blocks(document, result, page_image, section)

    body_children = list(document.element.body)
    assert [child.tag for child in body_children[:-1]] == [
        qn("w:p"),
        qn("w:p"),
        qn("w:tbl"),
        qn("w:p"),
    ]
    assert [paragraph.text for paragraph in document.paragraphs[:2]] == [
        "Document title",
        "Opening paragraph",
    ]
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.tables[0].cell(0, 0).text == "Cell"
    assert list(body_children[3].iter(qn("w:drawing")))


def test_reading_order_keeps_columns_together():
    blocks = [
        {"text": "right 1", "bbox": [220, 80, 380, 110]},
        {"text": "left 1", "bbox": [20, 80, 180, 110]},
        {"text": "right 2", "bbox": [220, 130, 380, 160]},
        {"text": "left 2", "bbox": [20, 130, 180, 160]},
    ]

    ordered = order_positioned_items(blocks, lambda item: item["bbox"], 400)

    assert [item["text"] for item in ordered] == [
        "left 1",
        "left 2",
        "right 1",
        "right 2",
    ]


def test_digital_text_retains_span_formatting():
    page = SimpleNamespace(
        rect=SimpleNamespace(width=600),
        get_text=lambda _kind: {
            "blocks": [
                {
                    "type": 0,
                    "bbox": [60, 40, 540, 80],
                    "lines": [
                        {
                            "spans": [
                                {
                                    "text": "Styled",
                                    "font": "ABCDEF+Arial",
                                    "size": 14,
                                    "flags": 18,
                                    "color": 0x123456,
                                }
                            ]
                        }
                    ],
                }
            ]
        },
    )
    document = Document()

    add_digital_page_text(document, page)

    run = document.paragraphs[0].runs[0]
    assert run.text == "Styled"
    assert run.bold is True
    assert run.italic is True
    assert run.font.name == "Arial"
    assert run.font.size.pt == 14
    assert str(run.font.color.rgb) == "123456"


def test_ocr_postprocessing_joins_only_line_break_hyphenation():
    assert normalized_text_lines("μετα-\nτροπή\nPDF-123") == ["μετατροπή", "PDF-123"]
