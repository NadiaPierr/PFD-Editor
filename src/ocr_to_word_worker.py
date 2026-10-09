"""OCR a scanned PDF into a Word document outside the GUI process."""

import os
import json
from pathlib import Path
import re
import sys
import tempfile
import traceback
from html.parser import HTMLParser
from io import BytesIO
from collections.abc import Mapping

import fitz
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageEnhance, ImageOps


# Text layers produced by PDF applications are much more accurate than OCR.
# Keep the threshold deliberately modest: a page with only a page number or a
# few accessibility labels should still be treated as a scan.
MINIMUM_DIGITAL_TEXT_CHARACTERS = 40
MINIMUM_DIGITAL_ALPHANUMERIC_CHARACTERS = 12
OCR_TEXT_DETECTION_LIMIT = 1920


def configure_gpu_dlls():
    """Make CUDA DLLs installed in this virtual environment visible on Windows."""
    site_packages = Path(sys.executable).resolve().parent.parent / "Lib" / "site-packages"
    dll_directories = [
        site_packages / "nvidia" / "cu13" / "bin" / "x86_64",
        site_packages / "nvidia" / "cudnn",
    ]
    handles = []
    existing_path = os.environ.get("PATH", "")
    for directory in dll_directories:
        if directory.is_dir():
            handles.append(os.add_dll_directory(str(directory)))
            existing_path = f"{directory}{os.pathsep}{existing_path}"
    os.environ["PATH"] = existing_path
    return handles


class TableHTMLParser(HTMLParser):
    """Parse Paddle table HTML, including merged cells and line breaks."""

    def __init__(self):
        super().__init__()
        self.rows = []
        self.current_row = None
        self.current_cell = None

    def handle_starttag(self, tag, _attrs):
        attributes = dict(_attrs)
        if tag == "tr":
            self.current_row = []
        elif tag in {"td", "th"} and self.current_row is not None:
            self.current_cell = {
                "parts": [],
                "rowspan": max(1, int(attributes.get("rowspan", "1") or 1)),
                "colspan": max(1, int(attributes.get("colspan", "1") or 1)),
                "header": tag == "th",
            }
        elif tag == "br" and self.current_cell is not None:
            self.current_cell["parts"].append("\n")

    def handle_data(self, data):
        if self.current_cell is not None:
            self.current_cell["parts"].append(data)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.current_cell is not None:
            self.current_cell["text"] = "".join(self.current_cell.pop("parts")).strip()
            self.current_row.append(self.current_cell)
            self.current_cell = None
        elif tag == "tr" and self.current_row:
            self.rows.append(self.current_row)
            self.current_row = None


def add_html_table(document, html):
    parser = TableHTMLParser()
    parser.feed(html)
    rows = [row for row in parser.rows if row]
    if not rows:
        document.add_paragraph(html)
        return

    # Lay cells into a rectangular grid first, so Word merges can be applied
    # correctly even when an earlier cell spans rows or columns.
    occupied = {}
    placements = []
    widest_column = 0
    for row_index, row in enumerate(rows):
        column_index = 0
        for cell in row:
            while (row_index, column_index) in occupied:
                column_index += 1
            rowspan, colspan = cell["rowspan"], cell["colspan"]
            for grid_row in range(row_index, row_index + rowspan):
                for grid_column in range(column_index, column_index + colspan):
                    occupied[(grid_row, grid_column)] = cell
            placements.append((row_index, column_index, rowspan, colspan, cell))
            widest_column = max(widest_column, column_index + colspan)
            column_index += colspan

    row_count = max((row + rowspan for row, _column, rowspan, _colspan, _cell in placements), default=0)
    if not row_count or not widest_column:
        document.add_paragraph(html)
        return
    table = document.add_table(rows=row_count, cols=widest_column)
    table.style = "Table Grid"
    for row_index, column_index, rowspan, colspan, cell_data in placements:
        anchor = table.cell(row_index, column_index)
        if rowspan > 1 or colspan > 1:
            anchor = anchor.merge(table.cell(row_index + rowspan - 1, column_index + colspan - 1))
        paragraph = anchor.paragraphs[0]
        for line_index, line in enumerate(cell_data["text"].splitlines() or [""]):
            if line_index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            run.bold = cell_data["header"]


def configure_section_from_pdf_page(section, page):
    """Make each Word section match its source PDF page geometry."""
    section.page_width = Inches(page.rect.width / 72)
    section.page_height = Inches(page.rect.height / 72)
    # Small, consistent margins give content room while respecting the page.
    margin = min(Inches(0.5), section.page_width // 12, section.page_height // 12)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin


def available_content_width(section):
    return section.page_width - section.left_margin - section.right_margin


def add_image_block(document, page_image, bbox, section):
    if bbox is None or len(bbox) != 4:
        return
    left, top, right, bottom = (int(value) for value in bbox)
    left = max(0, left)
    top = max(0, top)
    right = min(page_image.width, right)
    bottom = min(page_image.height, bottom)
    if right <= left or bottom <= top:
        return

    crop = page_image.crop((left, top, right, bottom))
    image_bytes = BytesIO()
    crop.save(image_bytes, format="PNG")
    image_bytes.seek(0)
    # Scale according to the visual width on the original page, preserving
    # aspect ratio and never exceeding the usable Word page width.
    source_width = max(page_image.width, 1)
    visual_width = available_content_width(section) * (right - left) / source_width
    width = max(Inches(0.35), min(visual_width, available_content_width(section)))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(image_bytes, width=width)


def normalized_text_lines(content):
    """Clean OCR lines and join words split only by a visual line break."""
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    normalized = []
    for line in lines:
        if (
            normalized
            and normalized[-1].endswith(("-", "‐", "‑"))
            and line
            and line[0].islower()
        ):
            normalized[-1] = normalized[-1][:-1] + line
        else:
            normalized.append(line)
    return normalized


def add_text_block(document, content, label):
    lines = normalized_text_lines(content)
    if not lines:
        return
    if "title" in label:
        level = 1 if label in {"title", "doc_title"} else 2
        document.add_heading(" ".join(lines), level=level)
        return
    # Layout blocks normally represent paragraphs.  Keeping their internal
    # line breaks avoids fragmenting a single paragraph into many Word ones.
    paragraph = document.add_paragraph()
    if label in {"header", "footer", "page_number"}:
        paragraph.style = "Caption"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    for line_index, line in enumerate(lines):
        if line_index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)


def page_has_usable_text_layer(page):
    """Return whether the PDF page has real, extractable text.

    Many scanned PDFs contain a tiny invisible text fragment (often a page
    number).  Requiring both a reasonable amount of text and alphanumeric
    characters prevents those files from accidentally bypassing OCR.
    """
    text = page.get_text("text").strip()
    alphanumeric_count = sum(character.isalnum() for character in text)
    return (
        len(text) >= MINIMUM_DIGITAL_TEXT_CHARACTERS
        and alphanumeric_count >= MINIMUM_DIGITAL_ALPHANUMERIC_CHARACTERS
    )


def pdf_color(value):
    """Convert PyMuPDF's packed 0xRRGGBB colour to a python-docx colour."""
    value = int(value or 0)
    return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)


def add_digital_page_text(document, page, section=None):
    """Copy editable PDF text while retaining span-level visual formatting."""
    section = section or document.sections[-1]
    page_width = max(float(page.rect.width), 1.0)
    usable_width = available_content_width(section)
    blocks = [
        block
        for block in page.get_text("dict").get("blocks", [])
        if block.get("type") == 0 and block.get("lines")
    ]

    for block in order_positioned_items(blocks, lambda item: item.get("bbox"), page_width):
        paragraph = document.add_paragraph()
        left, _top, right, _bottom = block.get("bbox", (0, 0, page_width, 0))
        paragraph.paragraph_format.left_indent = int(usable_width * max(0, left) / page_width)
        paragraph.paragraph_format.right_indent = int(
            usable_width * max(0, page_width - right) / page_width
        )
        paragraph.paragraph_format.space_after = Pt(3)

        for line_index, line in enumerate(block.get("lines", [])):
            if line_index:
                paragraph.add_run().add_break()
            spans = line.get("spans", [])
            for span in spans:
                text = span.get("text", "")
                if not text:
                    continue
                run = paragraph.add_run(text)
                flags = int(span.get("flags", 0))
                run.bold = bool(flags & 16)
                run.italic = bool(flags & 2)
                run.font.superscript = bool(flags & 1)
                run.font.size = Pt(max(1.0, float(span.get("size", 11))))
                run.font.color.rgb = pdf_color(span.get("color", 0))
                font_name = span.get("font")
                if font_name:
                    run.font.name = str(font_name).split("+")[-1]


def render_scale_for_page(page):
    """Choose a safe OCR scale from the actual PDF page dimensions.

    The former fixed 2.5x scale wastes GPU memory on large pages and can be
    too low for small receipts.  Keep the long side around 2,400 pixels while
    ensuring the short side has enough detail for small printed characters.
    """
    short_side = max(min(page.rect.width, page.rect.height), 1.0)
    long_side = max(page.rect.width, page.rect.height, 1.0)
    scale_for_detail = 1650 / short_side
    scale_for_memory = 2400 / long_side
    return max(1.75, min(3.0, max(scale_for_detail, min(scale_for_memory, 3.0))))


def prepare_scanned_page_image(page):
    """Render and lightly normalize a scanned page before OCR.

    Autocontrast and a restrained contrast increase improve faded scans
    without binarising photographs, coloured seals, or table shading.
    PyMuPDF already applies a page's PDF rotation when rendering.
    """
    scale = render_scale_for_page(page)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    page_image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    page_image = ImageOps.exif_transpose(page_image)
    page_image = ImageOps.autocontrast(page_image, cutoff=1)
    return ImageEnhance.Contrast(page_image).enhance(1.08)


def order_positioned_items(items, bbox_getter, page_width):
    """Return a reading order using recursive whitespace (XY-cut) analysis.

    Unlike a plain top/left sort, this keeps the full left column together
    before moving to the right column, while full-width titles and footers are
    separated naturally by horizontal whitespace.
    """
    positioned = [item for item in items if bbox_getter(item) and len(bbox_getter(item)) == 4]
    unpositioned = [item for item in items if item not in positioned]

    def axis_gap(group, axis):
        starts = sorted(group, key=lambda item: float(bbox_getter(item)[axis]))
        best = None
        running_end = float(bbox_getter(starts[0])[axis + 2])
        for index in range(1, len(starts)):
            start = float(bbox_getter(starts[index])[axis])
            gap = start - running_end
            if gap > 0 and (best is None or gap > best[0]):
                best = (gap, index, starts)
            running_end = max(running_end, float(bbox_getter(starts[index])[axis + 2]))
        return best

    def xy_cut(group):
        if len(group) < 2:
            return group
        horizontal = axis_gap(group, 1)
        vertical = axis_gap(group, 0)
        # A vertical separator is decisive evidence of columns. Full-width
        # titles/footers overlap both columns, so they naturally prevent that
        # cut until a horizontal band has separated them from the body.
        cut = vertical or horizontal
        if cut is None:
            return sorted(group, key=lambda item: (bbox_getter(item)[1], bbox_getter(item)[0]))
        _gap, index, ordered = cut
        return xy_cut(ordered[:index]) + xy_cut(ordered[index:])

    return xy_cut(positioned) + unpositioned


def add_layout_blocks(document, result, page_image, section):
    blocks = result.get("parsing_res_list", [])
    normalized_blocks = []
    for block in blocks:
        # Live PaddleX results contain LayoutBlock objects; exported JSON uses
        # block_* keys. Support both without losing the live object's fields.
        if isinstance(block, Mapping):
            label = block.get("block_label", block.get("label", "text"))
            content = block.get("block_content", block.get("content", ""))
            bbox = block.get("block_bbox", block.get("bbox"))
        else:
            label = block.label
            content = block.content
            bbox = block.bbox
        label = str(label).lower()
        content = str(content or "").strip()
        normalized_blocks.append((label, content, bbox))

    ordered_blocks = order_positioned_items(
        normalized_blocks, lambda item: item[2], page_image.width
    )
    for label, content, bbox in ordered_blocks:
        if label == "table":
            add_html_table(document, content)
        elif label in {"image", "figure", "chart"}:
            add_image_block(document, page_image, bbox, section)
        else:
            add_text_block(document, content, label)


def report_progress(status_path, message):
    """Append a structured progress event that the GUI can safely consume."""
    event = {"type": "progress", "message": message}
    with Path(status_path).open("a", encoding="utf-8") as status_file:
        status_file.write(json.dumps(event, ensure_ascii=False) + "\n")
        status_file.flush()


def acquire_conversion_lock():
    """Allow only one memory-intensive PP-Structure conversion at a time."""
    import msvcrt

    lock_path = Path(tempfile.gettempdir()) / "pdf_editor_scan_to_word.lock"
    handle = open(lock_path, "a+b")
    try:
        handle.write(b"0")
        handle.flush()
        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
    except OSError as exc:
        handle.close()
        raise RuntimeError(
            "Another Scan PDF to Word conversion is already running. "
            "Wait for it to finish, then try again."
        ) from exc
    return handle


class DownloadProgressStream:
    """Relays Paddle's model-download progress to the GUI status file."""

    def __init__(self, stream, status_path):
        self.stream = stream
        self.status_path = status_path
        self.model_name = "OCR model"
        self.last_message = ""

    def write(self, data):
        self.stream.write(data)
        model_match = re.search(r"Using official model \(([^)]+)\)", data)
        if model_match:
            self.model_name = model_match.group(1)
        percentage_match = re.search(r"(\d+(?:\.\d+)?)%", data)
        if percentage_match:
            message = f"Downloading {self.model_name}: {percentage_match.group(1)}%..."
            if message != self.last_message:
                report_progress(self.status_path, message)
                self.last_message = message
        return len(data)

    def flush(self):
        self.stream.flush()

    def __getattr__(self, name):
        return getattr(self.stream, name)


def convert(pdf_path, output_path, language, status_path):
    with acquire_conversion_lock():
        document = Document()
        with fitz.open(pdf_path) as pdf_document:
            native_text_pages = [
                page_has_usable_text_layer(pdf_document.load_page(page_index))
                for page_index in range(pdf_document.page_count)
            ]
            structure = None
            if any(not uses_native_text for uses_native_text in native_text_pages):
                report_progress(
                    status_path,
                    "Loading layout and table models for scanned pages (first use may take several minutes)...",
                )
                os.environ.setdefault("PADDLE_PDX_MODEL_SOURCE", "BOS")
                os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
                _dll_handles = configure_gpu_dlls()

                from paddleocr import PPStructureV3

                original_stdout, original_stderr = sys.stdout, sys.stderr
                sys.stdout = DownloadProgressStream(original_stdout, status_path)
                sys.stderr = DownloadProgressStream(original_stderr, status_path)
                try:
                    structure = PPStructureV3(
                        lang=language,
                        device="gpu:0",
                        use_doc_orientation_classify=True,
                        use_doc_unwarping=True,
                        use_textline_orientation=True,
                        use_table_recognition=True,
                        use_formula_recognition=False,
                        use_chart_recognition=False,
                        format_block_content=True,
                        text_det_limit_side_len=OCR_TEXT_DETECTION_LIMIT,
                    )
                finally:
                    sys.stdout, sys.stderr = original_stdout, original_stderr
            else:
                report_progress(status_path, "Using text already embedded in the PDF; OCR is not needed.")

            for page_index in range(pdf_document.page_count):
                page = pdf_document.load_page(page_index)
                page_number = page_index + 1
                total_pages = pdf_document.page_count
                # A Word section is the closest editable equivalent of a PDF
                # page: it lets mixed portrait/landscape PDFs retain their
                # original geometry without rasterising the entire page.
                section = (
                    document.sections[0]
                    if page_index == 0
                    else document.add_section(WD_SECTION.NEW_PAGE)
                )
                configure_section_from_pdf_page(section, page)
                if native_text_pages[page_index]:
                    report_progress(
                        status_path,
                        f"Extracting digital text from page {page_number} of {total_pages}...",
                    )
                    add_digital_page_text(document, page, section)
                else:
                    if structure is None:
                        raise RuntimeError("The OCR engine was not initialized for a scanned page.")
                    render_scale = render_scale_for_page(page)
                    report_progress(
                        status_path,
                        f"Analyzing scanned page {page_number} of {total_pages} at {render_scale:.1f}x...",
                    )
                    page_image = prepare_scanned_page_image(page)
                    results = structure.predict(np.asarray(page_image))
                    for result in results:
                        add_layout_blocks(document, result, page_image, section)
        report_progress(status_path, "Creating Word document...")
        output_path = Path(output_path)
        if not output_path.parent.is_dir():
            raise RuntimeError("The selected output folder does not exist.")

        # Never write directly to the user's destination.  A native-library
        # crash or a cancelled conversion can otherwise leave a corrupt DOCX.
        temporary_output = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{output_path.stem}_",
                suffix=".docx",
                dir=output_path.parent,
                delete=False,
            ) as output_file:
                temporary_output = Path(output_file.name)
            document.save(temporary_output)
            if not temporary_output.is_file() or temporary_output.stat().st_size == 0:
                raise RuntimeError("Word document creation did not produce an output file.")
            os.replace(temporary_output, output_path)
            temporary_output = None
        finally:
            if temporary_output is not None:
                try:
                    temporary_output.unlink(missing_ok=True)
                except OSError:
                    pass


if __name__ == "__main__":
    error_destination = Path(sys.argv[4])
    try:
        convert(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[5])
    except Exception as exc:
        error_destination.write_text(
            json.dumps(
                {
                    "message": f"OCR conversion failed: {exc}",
                    "traceback": traceback.format_exc(),
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        raise
